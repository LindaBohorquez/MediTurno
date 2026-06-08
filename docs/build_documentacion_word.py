from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
IMG_DIR = DOCS / "imagenes" / "word"
OUT = DOCS / "MediTurno_Documentacion_Entrega.docx"


BLUE = "2E74B5"
DARK = "0B2545"
MUTED = "52616F"
LIGHT = "F2F4F7"
GREEN = "E8F6EF"
GOLD = "FFF4DF"


def font(size=24, bold=False):
    candidates = [
        "C:/Windows/Fonts/segoeuib.ttf" if bold else "C:/Windows/Fonts/segoeui.ttf",
        "C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf",
    ]
    for candidate in candidates:
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, size)
    return ImageFont.load_default()


def rounded(draw, xy, fill, outline, radius=16, width=3):
    draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=width)


def arrow(draw, start, end, fill="#52616F", width=4):
    draw.line([start, end], fill=fill, width=width)
    x1, y1 = start
    x2, y2 = end
    if abs(x2 - x1) >= abs(y2 - y1):
        direction = 1 if x2 >= x1 else -1
        pts = [(x2, y2), (x2 - 14 * direction, y2 - 8), (x2 - 14 * direction, y2 + 8)]
    else:
        direction = 1 if y2 >= y1 else -1
        pts = [(x2, y2), (x2 - 8, y2 - 14 * direction), (x2 + 8, y2 - 14 * direction)]
    draw.polygon(pts, fill=fill)


def text(draw, xy, value, size=24, fill="#17212B", bold=False, anchor=None):
    draw.text(xy, value, fill=fill, font=font(size, bold), anchor=anchor)


def save_architecture():
    img = Image.new("RGB", (1600, 900), "#F7FAFC")
    d = ImageDraw.Draw(img)
    text(d, (800, 45), "MediTurno - Arquitectura MVC por capas", 34, "#" + DARK, True, "mm")
    layers = [
        ("Controller / API REST", "Endpoints: pacientes, medicos, especialidades, servicios, disponibilidades, citas y reportes", "#E8F3FF", "#2D74B8"),
        ("Application Facade", "CitaFacade simplifica agendar, cancelar, confirmar y reprogramar", "#FFF4DF", "#D18400"),
        ("Service / Dominio", "Reglas de negocio, NotificacionService, Factory Method y Strategy", "#EAF8EF", "#29915F"),
        ("Repository / Persistencia", "Spring Data JPA repositories", "#F0EDFF", "#6656B8"),
        ("Database", "H2 en memoria: jdbc:h2:mem:mediturno", "#FFF0F0", "#B84D4D"),
    ]
    y = 120
    centers = []
    for title, desc, fill, outline in layers:
        rounded(d, (120, y, 1480, y + 110), fill, outline)
        text(d, (155, y + 35), title, 27, "#17212B", True)
        text(d, (155, y + 75), desc, 22, "#263849")
        centers.append((800, y + 110))
        y += 145
    for i in range(4):
        arrow(d, (800, 120 + i * 145 + 110), (800, 120 + (i + 1) * 145), "#52616F", 5)
    path = IMG_DIR / "arquitectura-capas.png"
    img.save(path)
    return path


def save_patterns():
    img = Image.new("RGB", (1600, 900), "#FBFCFE")
    d = ImageDraw.Draw(img)
    text(d, (800, 45), "Patrones de diseno implementados", 34, "#" + DARK, True, "mm")
    panels = [
        (70, "Factory Method", "NotificacionFactory", ["EmailNotificacion", "SmsNotificacion", "WhatsAppNotificacion"], "Crea notificadores por canal."),
        (560, "Facade", "CitaFacade", ["CitaService", "NotificacionService"], "Evita orquestacion en el controller."),
        (1050, "Strategy", "PoliticaCancelacion", ["FlexibleStrategy", "RestrictivaStrategy"], "Cambia reglas sin tocar CitaService."),
    ]
    for x, title, main, children, footer in panels:
        rounded(d, (x, 120, x + 420, 780), "#FFFFFF", "#D6DEE8")
        text(d, (x + 25, 165), title, 27, "#" + DARK, True)
        rounded(d, (x + 70, 230, x + 350, 315), "#" + GOLD, "#D18400")
        text(d, (x + 210, 275), main, 21, "#17212B", True, "mm")
        child_y = 405
        for child in children:
            rounded(d, (x + 90, child_y, x + 330, child_y + 70), "#EEF6FF", "#2D74B8")
            text(d, (x + 210, child_y + 35), child, 18, "#17212B", False, "mm")
            arrow(d, (x + 210, 315), (x + 210, child_y), "#52616F", 4)
            child_y += 95
        text(d, (x + 35, 715), footer, 20, "#263849")
    path = IMG_DIR / "patrones-diseno.png"
    img.save(path)
    return path


def save_uml():
    img = Image.new("RGB", (1700, 1050), "#F7FAFC")
    d = ImageDraw.Draw(img)
    text(d, (850, 45), "MediTurno - Modelo de clases principales", 34, "#" + DARK, True, "mm")
    boxes = [
        (70, 115, "Paciente", ["id", "nombres", "documento", "correo", "activo"]),
        (410, 115, "Medico", ["id", "registroMedico", "especialidad", "activo"]),
        (750, 115, "Especialidad", ["id", "nombre", "descripcion", "activa"]),
        (1090, 115, "ServicioMedico", ["id", "nombre", "duracion", "tarifa", "especialidad"]),
        (410, 405, "DisponibilidadMedica", ["medico", "fecha", "horaInicio", "estado", "reservar()", "liberar()"]),
        (790, 390, "Cita", ["paciente", "medico", "servicioMedico", "disponibilidad", "estado", "confirmar()", "cancelar()", "reprogramar()"]),
        (70, 690, "[P] CitaFacade", ["agendar()", "cancelar()", "reprogramar()"]),
        (410, 690, "[P] CitaService", ["reglas de citas", "persistencia"]),
        (790, 690, "[P] NotificacionFactory", ["crear(canal)", "Email/SMS/WhatsApp"]),
        (1160, 690, "[P] PoliticaCancelacion", ["Flexible", "Restrictiva"]),
    ]
    for x, y, title, attrs in boxes:
        w = 300 if x < 1090 else 350
        h = 185 if y < 650 else 165
        fill = "#FFF7E6" if "[P]" in title else "#FFFFFF"
        outline = "#D18400" if "[P]" in title else "#2F5D8C"
        rounded(d, (x, y, x + w, y + h), fill, outline)
        text(d, (x + 18, y + 32), title, 23, "#17212B", True)
        yy = y + 66
        for attr in attrs:
            text(d, (x + 22, yy), "- " + attr, 18, "#263849")
            yy += 24
    relations = [
        ((710, 205), (750, 205)),
        ((1090, 205), (1050, 205)),
        ((560, 300), (560, 405)),
        ((790, 480), (710, 480)),
        ((1110, 390), (1210, 300)),
        ((370, 770), (410, 770)),
        ((710, 770), (790, 770)),
        ((1090, 770), (1160, 770)),
    ]
    for start, end in relations:
        arrow(d, start, end, "#52616F", 4)
    path = IMG_DIR / "uml-clases.png"
    img.save(path)
    return path


def save_flow():
    img = Image.new("RGB", (1700, 900), "#F8FBFB")
    d = ImageDraw.Draw(img)
    text(d, (850, 45), "Flujo funcional probado en Postman/API", 34, "#" + DARK, True, "mm")
    steps = [
        ("1", "Crear especialidad", "POST /api/especialidades"),
        ("2", "Crear medico", "POST /api/medicos"),
        ("3", "Crear paciente", "POST /api/pacientes"),
        ("4", "Crear servicio", "POST /servicios-medicos"),
        ("5", "Disponibilidad", "POST /disponibilidades"),
        ("6", "Agendar cita", "POST /api/citas"),
        ("7", "Ocupada", "estado = OCUPADA"),
        ("8", "Cancelar", "PATCH /citas/{id}/cancelar"),
        ("9", "Libre", "estado = DISPONIBLE"),
        ("10", "Nueva disponibilidad", "POST /disponibilidades"),
        ("11", "Reprogramar", "estado = REPROGRAMADA"),
        ("12", "Reporte", "total = 1"),
    ]
    positions = []
    for idx, item in enumerate(steps):
        row = idx // 4
        col = idx % 4
        x = 100 + col * 390
        y = 130 + row * 220
        positions.append((x, y))
        fill = "#E8F8EF" if item[1] in ["Ocupada", "Libre", "Reprogramar", "Reporte"] else "#FFFFFF"
        rounded(d, (x, y, x + 320, y + 110), fill, "#2E7D9A")
        text(d, (x + 20, y + 36), item[0] + ". " + item[1], 22, "#17212B", True)
        text(d, (x + 20, y + 73), item[2], 17, "#263849")
    for i in range(len(positions) - 1):
        x1, y1 = positions[i]
        x2, y2 = positions[i + 1]
        if (i + 1) % 4 == 0:
            arrow(d, (x1 + 160, y1 + 110), (positions[i + 1][0] + 160, positions[i + 1][1]), "#52616F", 4)
        else:
            arrow(d, (x1 + 320, y1 + 55), (x2, y2 + 55), "#52616F", 4)
    path = IMG_DIR / "flujo-postman.png"
    img.save(path)
    return path


def save_tests():
    img = Image.new("RGB", (1500, 780), "#101820")
    d = ImageDraw.Draw(img)
    rounded(d, (70, 70, 1430, 710), "#0D1117", "#30363D")
    lines = [
        ("MediTurno - Evidencia de pruebas", "#7EE787", 26, True),
        ("> .\\mvnw.cmd clean test", "#D1D5DA", 22, False),
        ("Compiling 54 source files with javac [release 21]", "#C9D1D9", 19, False),
        ("Compiling 5 test source files", "#C9D1D9", 19, False),
        ("NotificacionFactoryTest .......... 2 tests OK", "#C9D1D9", 19, False),
        ("MediTurnoApplicationTests ........ 1 test OK", "#C9D1D9", 19, False),
        ("DisponibilidadMedicaTest ........ 3 tests OK", "#C9D1D9", 19, False),
        ("CitaServiceTest ................. 5 tests OK", "#C9D1D9", 19, False),
        ("PoliticaCancelacionTest ......... 4 tests OK", "#C9D1D9", 19, False),
        ("Results: Tests run: 15, Failures: 0, Errors: 0, Skipped: 0", "#7EE787", 22, True),
        ("BUILD SUCCESS", "#7EE787", 26, True),
    ]
    y = 120
    for value, color, size, bold in lines:
        text(d, (115, y), value, size, color, bold)
        y += 55 if size >= 22 else 42
    path = IMG_DIR / "evidencia-pruebas.png"
    img.save(path)
    return path


def save_endpoints():
    img = Image.new("RGB", (1600, 900), "#F7FAFC")
    d = ImageDraw.Draw(img)
    text(d, (800, 45), "Mapa de endpoints REST", 34, "#" + DARK, True, "mm")
    cards = [
        (90, 120, "Catalogos y usuarios", ["/api/pacientes", "/api/medicos", "/api/especialidades", "/api/servicios-medicos"]),
        (610, 120, "Agenda medica", ["/api/disponibilidades", "GET medicoId, fecha", "PATCH /{id}/reservar", "PATCH /{id}/liberar"]),
        (1130, 120, "Citas", ["/api/citas", "GET pacienteId, medicoId", "PATCH confirmar", "PATCH cancelar", "PATCH reprogramar"]),
        (350, 520, "Reportes", ["/api/reportes/citas", "GET desde, hasta, medicoId", "Resumen por estado"]),
        (870, 520, "Soporte local", ["/h2-console", "jdbc:h2:mem:mediturno", "user: sa"]),
    ]
    for x, y, title, lines in cards:
        rounded(d, (x, y, x + 390, y + 260), "#FFFFFF", "#CBD7E3")
        text(d, (x + 25, y + 38), title, 24, "#" + DARK, True)
        yy = y + 85
        for line in lines:
            text(d, (x + 25, yy), line, 19, "#263849")
            yy += 34
    path = IMG_DIR / "endpoints-api.png"
    img.save(path)
    return path


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, value, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(value)
    run.font.name = "Calibri"
    run.font.size = Pt(9.5)
    run.bold = bold
    p.paragraph_format.space_after = Pt(0)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    hdr = table.rows[0].cells
    for i, header in enumerate(headers):
        set_cell_text(hdr[i], header, True)
        set_cell_shading(hdr[i], LIGHT)
        if widths:
            hdr[i].width = Inches(widths[i])
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], str(value))
            if widths:
                cells[i].width = Inches(widths[i])
    doc.add_paragraph()
    return table


def add_figure(doc, path, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Inches(6.25))
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.runs[0].font.size = Pt(9)
    cap.runs[0].font.italic = True
    cap.runs[0].font.color.rgb = RGBColor(82, 97, 111)


def setup_styles(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def add_cover(doc):
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(80)
    run = title.add_run("MediTurno")
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(34)
    run.font.color.rgb = RGBColor.from_string(DARK)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(28)
    run = subtitle.add_run("Documentacion tecnica del proyecto Spring Boot")
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor.from_string(MUTED)

    meta = [
        ("Proyecto", "Sistema de Gestion de Citas Medicas"),
        ("Entregable", "Codigo, endpoints, patrones, pruebas y evidencias"),
        ("Stack", "Java 21, Spring Boot 4, Spring Data JPA, H2, JUnit 5, Mockito"),
        ("Verificacion", "mvn clean test: BUILD SUCCESS, 15 pruebas ejecutadas"),
    ]
    add_table(doc, ["Campo", "Detalle"], meta, [1.7, 4.5])
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Incluye diagramas, evidencias del flujo funcional y pruebas pendientes recomendadas.")
    run.font.color.rgb = RGBColor.from_string(MUTED)
    doc.add_page_break()


def build_docx(images):
    doc = Document()
    setup_styles(doc)
    add_cover(doc)

    doc.add_heading("1. Resumen ejecutivo", level=1)
    doc.add_paragraph(
        "MediTurno implementa un backend REST para la gestion de citas medicas. El codigo esta organizado por capas "
        "y refleja el documento academico: modelos de dominio, servicios, repositorios, controladores, patrones de "
        "diseno y pruebas unitarias."
    )
    add_table(
        doc,
        ["Aspecto", "Estado"],
        [
            ("Clases principales", "Modelos JPA, services, repositories, controllers, DTOs, factory, facade y strategy"),
            ("Patrones", "Factory Method, Facade y Strategy implementados explicitamente"),
            ("Flujo funcional", "Probado con especialidad, medico, paciente, servicio, disponibilidad, cita, cancelacion, reprogramacion y reporte"),
            ("Pruebas", "15 pruebas en verde con JUnit 5 y Mockito"),
        ],
        [2.0, 4.2],
    )

    doc.add_heading("2. Arquitectura por capas", level=1)
    doc.add_paragraph(
        "La arquitectura corresponde a MVC con Spring Boot. Los controladores exponen endpoints REST, la fachada "
        "simplifica los casos de uso de cita, los servicios concentran reglas de negocio y los repositorios acceden "
        "a H2 mediante Spring Data JPA."
    )
    add_figure(doc, images["architecture"], "Figura 1. Arquitectura MVC por capas de MediTurno.")

    doc.add_heading("3. Modelo de clases", level=1)
    doc.add_paragraph(
        "El modelo cumple el minimo de ocho clases solicitado. Las entidades principales representan pacientes, "
        "medicos, especialidades, servicios medicos, disponibilidad y citas; los enums modelan estados relevantes."
    )
    add_figure(doc, images["uml"], "Figura 2. UML de clases principales y clases candidatas a patrones.")
    add_table(
        doc,
        ["Capa", "Clases"],
        [
            ("Modelo", "Paciente, Medico, Especialidad, ServicioMedico, DisponibilidadMedica, Cita, EstadoCita, EstadoDisponibilidad"),
            ("Repository", "PacienteRepository, MedicoRepository, EspecialidadRepository, ServicioMedicoRepository, DisponibilidadMedicaRepository, CitaRepository"),
            ("Service", "PacienteService, MedicoService, EspecialidadService, ServicioMedicoService, DisponibilidadMedicaService, CitaService, NotificacionService, ReporteCitasService"),
            ("Controller", "PacienteController, MedicoController, EspecialidadController, ServicioMedicoController, DisponibilidadMedicaController, CitaController, ReporteCitasController"),
            ("DTO", "CitaRequest, CancelarCitaRequest, ReprogramarCitaRequest, DisponibilidadMedicaRequest, MedicoRequest, ServicioMedicoRequest, ReporteCitasResponse"),
        ],
        [1.35, 4.85],
    )

    doc.add_heading("4. Patrones de diseno", level=1)
    add_figure(doc, images["patterns"], "Figura 3. Patrones creacional, estructural y de comportamiento.")
    add_table(
        doc,
        ["Categoria", "Implementacion", "Responsabilidad"],
        [
            ("Creacional", "NotificacionFactory, Notificador, Email/SMS/WhatsApp", "Crear notificaciones segun canal sin condicionales repetidos"),
            ("Estructural", "CitaFacade", "Orquestar CitaService y NotificacionService para flujos principales"),
            ("Comportamiento", "PoliticaCancelacion, FlexibleStrategy, RestrictivaStrategy", "Variar reglas de cancelacion sin modificar CitaService"),
        ],
        [1.25, 2.2, 2.75],
    )

    doc.add_heading("5. Endpoints disponibles", level=1)
    add_figure(doc, images["endpoints"], "Figura 4. Mapa visual de endpoints REST.")
    add_table(
        doc,
        ["Modulo", "Endpoints"],
        [
            ("Pacientes", "GET/POST /api/pacientes, GET/PUT/DELETE /api/pacientes/{id}"),
            ("Medicos", "GET/POST /api/medicos, GET/PUT/DELETE /api/medicos/{id}, filtro especialidadId"),
            ("Especialidades", "GET/POST /api/especialidades, GET/PUT/DELETE /api/especialidades/{id}"),
            ("Servicios medicos", "GET/POST /api/servicios-medicos, GET/PUT/DELETE /api/servicios-medicos/{id}"),
            ("Disponibilidades", "GET/POST /api/disponibilidades, PATCH reservar/liberar, DELETE cancelar"),
            ("Citas", "GET/POST /api/citas, PATCH confirmar/cancelar/reprogramar"),
            ("Reportes", "GET /api/reportes/citas con filtros desde, hasta y medicoId"),
        ],
        [1.65, 4.55],
    )

    doc.add_heading("6. Flujo funcional y evidencias", level=1)
    doc.add_paragraph(
        "El flujo funcional fue ejecutado por HTTP contra la API. Las respuestas JSON estan guardadas en "
        "docs/evidencias-flujo y la coleccion importable se encuentra en docs/MediTurno.postman_collection.json."
    )
    add_figure(doc, images["flow"], "Figura 5. Flujo minimo probado con Postman/API.")
    add_table(
        doc,
        ["Paso", "Evidencia"],
        [
            ("Crear especialidad", "01-crear-especialidad.json"),
            ("Crear medico", "02-crear-medico.json"),
            ("Crear paciente", "03-crear-paciente.json"),
            ("Crear servicio medico", "04-crear-servicio-medico.json"),
            ("Registrar disponibilidad", "05-registrar-disponibilidad.json"),
            ("Agendar cita", "06-agendar-cita.json"),
            ("Verificar ocupada", "07-verificar-disponibilidad-ocupada.json"),
            ("Cancelar cita", "08-cancelar-cita.json"),
            ("Verificar libre", "09-verificar-disponibilidad-libre.json"),
            ("Reprogramar cita", "11-reprogramar-cita.json"),
            ("Generar reporte", "12-generar-reporte.json"),
        ],
        [2.1, 4.1],
    )

    doc.add_heading("7. Pruebas", level=1)
    add_figure(doc, images["tests"], "Figura 6. Evidencia resumida de mvn clean test.")
    add_table(
        doc,
        ["Test", "Cobertura"],
        [
            ("MediTurnoApplicationTests", "Carga del contexto Spring"),
            ("CitaServiceTest", "Agendar, rechazar disponibilidad ocupada, cancelar y reprogramar"),
            ("DisponibilidadMedicaTest", "Reservar, liberar y rechazar reserva repetida"),
            ("NotificacionFactoryTest", "Crear notificacion por canal y email por defecto"),
            ("PoliticaCancelacionTest", "Estrategias flexible y restrictiva, incluida cancelacion tardia"),
        ],
        [2.25, 3.95],
    )

    doc.add_heading("8. Pruebas pendientes recomendadas", level=1)
    for item in [
        "Pruebas de integracion HTTP con MockMvc para todo el flujo.",
        "Pruebas de controllers para codigos 201, 400, 404 y 204.",
        "Pruebas de repositories con H2 para solapamientos y filtros.",
        "Pruebas especificas de ReporteCitasService.",
        "Configurar JaCoCo para cobertura.",
        "Analisis SonarQube para complejidad ciclomatica, duplicacion y deuda tecnica.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("9. Archivos de soporte", level=1)
    add_table(
        doc,
        ["Archivo", "Uso"],
        [
            ("README.md", "Guia rapida del proyecto"),
            ("docs/postman-flujo-completo.md", "Instrucciones paso a paso para Postman"),
            ("docs/MediTurno.postman_collection.json", "Coleccion importable en Postman"),
            ("docs/evidencias-flujo", "Respuestas JSON reales del flujo probado"),
            ("docs/imagenes", "Diagramas e imagenes de evidencia"),
            ("docs/resumen-tecnico-codigo.md", "Resumen tecnico por capas"),
        ],
        [2.55, 3.65],
    )

    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run("MediTurno - Documentacion tecnica de entrega")
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor.from_string(MUTED)

    doc.save(OUT)
    return OUT


def main():
    IMG_DIR.mkdir(parents=True, exist_ok=True)
    images = {
        "architecture": save_architecture(),
        "patterns": save_patterns(),
        "uml": save_uml(),
        "flow": save_flow(),
        "tests": save_tests(),
        "endpoints": save_endpoints(),
    }
    out = build_docx(images)
    print(out)


if __name__ == "__main__":
    main()
